#!/usr/bin/env python3
"""Сборка короткой презентации к защите: slides.json + screens/ → PPTX (+ PDF и PNG через LibreOffice).

    tools/heavy.sh .venv/bin/python deliverables/presentation-v2/build.py            только PPTX
    tools/heavy.sh .venv/bin/python deliverables/presentation-v2/build.py --pdf      PPTX и PDF
    tools/heavy.sh .venv/bin/python deliverables/presentation-v2/build.py --pdf --png /tmp/v2-png

Все элементы нативные: текстовые блоки, фигуры и картинки можно править в PowerPoint или Keynote.
Цвета — из frontend/src/styles/global.css, шрифт — Onest (если не установлен, подставится системный).
"""

from __future__ import annotations

import argparse
import json
import os
import shutil
import subprocess
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
ASSETS = HERE / "assets"

# Палитра интерфейса (frontend/src/styles/global.css).
BG0 = "07080a"
BG1 = "0c0e11"
BG2 = "111317"
BG3 = "171a1f"
BG4 = "1e2228"
LINE = "2c3037"
TEXT1 = "edebe6"
TEXT2 = "a9a7a0"
TEXT3 = "77756f"
TEXT4 = "504f4b"
BRAND = "f5c400"
BRAND_DIM = "2a2307"

FONT = "Onest"
FONT_MEDIUM = "Onest Medium"
FONT_SEMI = "Onest SemiBold"
FONT_XBOLD = "Onest ExtraBold"

SLIDE_W = 13.333
SLIDE_H = 7.5
MARGIN = 0.6
CONTENT_TOP = 1.95


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value.upper())


# ---------------------------------------------------------------- ассеты


def _font_file(name: str) -> Path | None:
    for base in (Path.home() / "Library/Fonts", Path("/Library/Fonts")):
        candidate = base / name
        if candidate.exists():
            return candidate
    return None


def make_hazard(path: Path) -> None:
    """Жёлто-чёрная штриховка — фирменный знак интерфейса."""
    w, h, step = 2400, 44, 44
    img = Image.new("RGB", (w, h), "#" + BG0)
    draw = ImageDraw.Draw(img)
    for x in range(-h, w + h, step):
        draw.polygon([(x, h), (x + step // 2, h), (x + step // 2 + h, 0), (x + h, 0)], fill="#" + BRAND)
    img.save(path)


def make_logo(path: Path) -> None:
    """Знак как в favicon интерфейса: штриховая рамка, «З» и её зеркальное отражение."""
    s = 600
    scale = 4
    big = s * scale
    img = Image.new("RGBA", (big, big), (0, 0, 0, 0))
    draw = ImageDraw.Draw(img)
    u = big / 32
    draw.rounded_rectangle([0, 0, big - 1, big - 1], radius=6 * u, fill="#111111")
    stripes = Image.new("RGBA", (big, big), "#111111")
    sd = ImageDraw.Draw(stripes)
    step = int(6 * u)
    for x in range(-big, 2 * big, step):
        sd.polygon([(x, big), (x + step // 2, big), (x + step // 2 + big, 0), (x + big, 0)], fill="#" + BRAND)
    mask = Image.new("L", (big, big), 0)
    ImageDraw.Draw(mask).rounded_rectangle([2 * u, 2 * u, 30 * u, 30 * u], radius=4 * u, fill=255)
    img.paste(stripes, (0, 0), mask)
    draw.rounded_rectangle([5 * u, 5 * u, 27 * u, 27 * u], radius=3 * u, fill="#0d0e10")
    font_path = _font_file("Onest-Bold.ttf")
    font = ImageFont.truetype(str(font_path), int(15 * u)) if font_path else ImageFont.load_default()
    glyph = Image.new("RGBA", (int(10 * u), int(18 * u)), (0, 0, 0, 0))
    gd = ImageDraw.Draw(glyph)
    gd.text((glyph.width / 2, glyph.height / 2), "З", font=font, fill="#" + BRAND, anchor="mm")
    img.alpha_composite(glyph, (int(6 * u), int(7 * u)))
    img.alpha_composite(glyph.transpose(Image.FLIP_LEFT_RIGHT), (int(16 * u), int(7 * u)))
    draw.rectangle([15.5 * u, 8 * u, 16.5 * u, 24 * u], fill="#" + BRAND)
    img.resize((s, s), Image.LANCZOS).save(path)


def ensure_assets() -> None:
    ASSETS.mkdir(exist_ok=True)
    if not (ASSETS / "hazard.png").exists():
        make_hazard(ASSETS / "hazard.png")
    if not (ASSETS / "logo.png").exists():
        make_logo(ASSETS / "logo.png")


# ---------------------------------------------------------------- примитивы


def set_background(slide, color: str = BG1) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def _strip_style(shape) -> None:
    """Убираем ссылку на стиль темы: без теней и чужих цветов."""
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)


def add_box(slide, x, y, w, h, fill=None, line=None, line_w=1.0, radius=0.08, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    box = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    _strip_style(box)
    if fill:
        box.fill.solid()
        box.fill.fore_color.rgb = rgb(fill)
    else:
        box.fill.background()
    if line:
        box.line.color.rgb = rgb(line)
        box.line.width = Pt(line_w)
    else:
        box.line.fill.background()
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE and radius is not None:
        box.adjustments[0] = min(0.5, radius / min(w, h))
    box.text_frame.text = ""
    return box


def style_run(run, size, color=TEXT1, font=FONT, bold=False, spacing=None) -> None:
    run.font.size = Pt(size)
    run.font.name = font
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    rpr = run._r.get_or_add_rPr()
    rpr.set("lang", "ru-RU")
    if spacing is not None:
        rpr.set("spc", str(int(spacing * 100)))
    for tag in ("a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = rpr.makeelement(qn(tag), {})
            rpr.append(el)
        el.set("typeface", font)


FONT_FILES = {
    FONT: "Onest-Regular.ttf",
    FONT_MEDIUM: "Onest-Medium.ttf",
    FONT_SEMI: "Onest-SemiBold.ttf",
    FONT_XBOLD: "Onest-ExtraBold.ttf",
}
_FONT_CACHE: dict = {}
WARNINGS: list[str] = []
CURRENT = {"slide": 0}
LINE_HEIGHT = 1.22  # высота строки Onest в долях кегля при множителе 1.0


def _measure_font(font: str, size: float):
    key = (font, size)
    if key not in _FONT_CACHE:
        path = _font_file(FONT_FILES.get(font, "Onest-Regular.ttf"))
        # Кегль ×10, чтобы округление пикселей не искажало ширину.
        _FONT_CACHE[key] = ImageFont.truetype(str(path), int(size * 10)) if path else None
    return _FONT_CACHE[key]


def count_lines(text: str, font: str, size: float, width_in: float, spacing: float = 0) -> int:
    """Сколько строк займёт текст при переносе по словам (ширина по метрикам Onest)."""
    measure = _measure_font(font, size)
    limit = width_in * 72 * 10
    total = 0
    for raw in str(text).split("\n"):
        words = raw.split(" ")
        lines, current = 1, ""
        for word in words:
            candidate = f"{current} {word}".strip()
            if measure is not None:
                width = measure.getlength(candidate) + spacing * 10 * len(candidate)
            else:
                width = len(candidate) * size * 0.55 * 10
            if width > limit and current:
                lines += 1
                current = word
            else:
                current = candidate
        total += lines
    return total


def text_height(paragraphs, width_in: float) -> float:
    total = 0.0
    for spec in paragraphs:
        size = spec.get("size", 18)
        lines = count_lines(spec["text"], spec.get("font", FONT), size, width_in, spec.get("spacing") or 0)
        total += lines * size * LINE_HEIGHT * spec.get("line", 1.05) / 72 + spec.get("before", 0) / 72
    return total


def add_text(slide, x, y, w, h, paragraphs, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, name=None):
    """paragraphs: [{text, size, color, font, bold, spacing, line, before}] — один абзац на элемент."""
    need = text_height(paragraphs, w)
    if need > h + 0.04:
        first = str(paragraphs[0]["text"]).replace("\n", " ")[:40]
        WARNINGS.append(f"слайд {CURRENT['slide']}: текст «{first}» выше блока на {need - h:.2f} дюйма")
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, spec in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = spec.get("align", align)
        p.line_spacing = spec.get("line", 1.05)
        if spec.get("before"):
            p.space_before = Pt(spec["before"])
        lines = str(spec["text"]).split("\n")
        for j, line in enumerate(lines):
            if j:
                p.add_line_break()
            run = p.add_run()
            run.text = line
            style_run(
                run,
                spec.get("size", 18),
                spec.get("color", TEXT1),
                spec.get("font", FONT),
                spec.get("bold", False),
                spec.get("spacing"),
            )
    return box


def image_aspect(path: Path, crop) -> float:
    with Image.open(path) as im:
        w, h = im.size
    cl, ct, cr, cb = crop
    return (w * (1 - cl - cr)) / (h * (1 - ct - cb))


def add_picture(slide, path: Path, x, y, w, h, crop=(0, 0, 0, 0), radius=0.07, border=LINE):
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    cl, ct, cr, cb = crop
    pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = cl, ct, cr, cb
    if radius:
        geom = pic._element.spPr.find(qn("a:prstGeom"))
        geom.set("prst", "roundRect")
        av = geom.find(qn("a:avLst"))
        if av is None:
            av = geom.makeelement(qn("a:avLst"), {})
            geom.append(av)
        gd = av.makeelement(qn("a:gd"), {"name": "adj", "fmla": f"val {int(radius / min(w, h) * 100000)}"})
        av.append(gd)
    if border:
        pic.line.color.rgb = rgb(border)
        pic.line.width = Pt(1)
    return pic


def fit(aspect, bx, by, bw, bh, halign="center", valign="top"):
    if bw / bh > aspect:
        h, w = bh, bh * aspect
    else:
        w, h = bw, bw / aspect
    x = {"left": bx, "center": bx + (bw - w) / 2, "right": bx + bw - w}[halign]
    y = {"top": by, "center": by + (bh - h) / 2, "bottom": by + bh - h}[valign]
    return x, y, w, h


def add_notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


def add_header(slide, spec, number: int, total: int) -> None:
    kicker = spec.get("kicker")
    if kicker:
        add_text(
            slide,
            MARGIN,
            0.48,
            8.0,
            0.3,
            [{"text": kicker.upper(), "size": 13, "color": BRAND, "font": FONT_SEMI, "spacing": 2}],
            name="Раздел",
        )
    add_text(
        slide,
        SLIDE_W - MARGIN - 2.0,
        0.48,
        2.0,
        0.3,
        [{"text": f"{number:02d} / {total:02d}", "size": 12, "color": TEXT4, "font": FONT_MEDIUM, "align": PP_ALIGN.RIGHT}],
        name="Номер",
    )
    add_text(
        slide,
        MARGIN,
        0.82,
        SLIDE_W - 2 * MARGIN,
        0.8,
        [{"text": spec["title"], "size": 36, "color": TEXT1, "font": FONT_XBOLD, "line": 1.0}],
        name="Заголовок",
    )


# ---------------------------------------------------------------- макеты


def layout_title(prs, spec, number, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, BG0)
    image = HERE / spec["image"]
    if image.exists():
        # Картинка уходит за правый край слайда: справа обрезаем ровно столько, сколько не помещается.
        crop = list(spec.get("crop", [0, 0, 0, 0]))
        img_x, img_h = 6.45, 5.0
        with Image.open(image) as im:
            pw, ph = im.size
        px_per_in = ph * (1 - crop[1] - crop[3]) / img_h
        visible_w = min(SLIDE_W - img_x, pw * (1 - crop[0] - crop[2]) / px_per_in)
        crop[2] = max(crop[2], 1 - crop[0] - visible_w * px_per_in / pw)
        add_picture(slide, image, SLIDE_W - visible_w, (SLIDE_H - img_h) / 2 - 0.05, visible_w, img_h, crop=crop, radius=0.12)
    logo = ASSETS / "logo.png"
    slide.shapes.add_picture(str(logo), Inches(MARGIN + 0.05), Inches(1.0), Inches(1.0), Inches(1.0))
    add_text(slide, MARGIN, 2.45, 5.4, 0.35, [{"text": spec["kicker"], "size": 16, "color": BRAND, "font": FONT_SEMI, "spacing": 1}])
    add_text(slide, MARGIN, 2.85, 5.6, 1.2, [{"text": spec["title"], "size": 66, "color": TEXT1, "font": FONT_XBOLD, "line": 0.95}])
    add_text(slide, MARGIN, 4.1, 5.6, 1.2, [{"text": spec["subtitle"], "size": 24, "color": TEXT1, "font": FONT_MEDIUM, "line": 1.1}])
    add_text(slide, MARGIN, 6.05, 5.4, 0.4, [{"text": spec["meta"], "size": 16, "color": TEXT2}])
    slide.shapes.add_picture(str(ASSETS / "hazard.png"), 0, Inches(SLIDE_H - 0.14), Inches(SLIDE_W), Inches(0.14))
    add_notes(slide, spec["notes"])


def layout_scheme(prs, spec, number, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_header(slide, spec, number, total)
    steps = spec["steps"]
    n = len(steps)
    gap = 0.55
    top = 2.2
    box_w = (SLIDE_W - 2 * MARGIN - gap * (n - 1)) / n
    pad = 0.28

    def step_paragraphs(step):
        return [
            {"text": step["title"], "size": 23, "color": TEXT1, "font": FONT_SEMI, "line": 1.0},
            {"text": step["text"], "size": 16, "color": TEXT2, "before": 6, "line": 1.1},
        ]

    text_h = max(text_height(step_paragraphs(step), box_w - 2 * pad) for step in steps)
    box_h = 0.62 + text_h + 0.3
    for i, step in enumerate(steps):
        x = MARGIN + i * (box_w + gap)
        accent = step.get("accent")
        add_box(slide, x, top, box_w, box_h, fill=BRAND_DIM if accent else BG3, line=BRAND if accent else LINE, line_w=1.5 if accent else 1.0, radius=0.14)
        add_text(slide, x + pad, top + 0.26, box_w - 2 * pad, 0.3, [{"text": f"{i + 1:02d}", "size": 14, "color": BRAND, "font": FONT_SEMI, "spacing": 1}])
        add_text(slide, x + pad, top + 0.62, box_w - 2 * pad, text_h + 0.05, step_paragraphs(step))
        if i < n - 1:
            ax = x + box_w + 0.12
            arrow = add_box(slide, ax, top + box_h / 2 - 0.17, gap - 0.28, 0.34, fill=BRAND, shape=MSO_SHAPE.RIGHT_ARROW, radius=None)
            arrow.name = "Стрелка"
    facts = spec.get("facts", [])
    if facts:
        fy = max(4.95, top + box_h + 0.75)
        col_w = (SLIDE_W - 2 * MARGIN) / len(facts)
        add_box(slide, MARGIN, fy - 0.35, SLIDE_W - 2 * MARGIN, 0.0, line=LINE, shape=MSO_SHAPE.RECTANGLE, radius=None)
        for i, fact in enumerate(facts):
            x = MARGIN + i * col_w
            add_text(
                slide,
                x,
                fy,
                col_w - 0.3,
                1.7,
                [
                    {"text": fact["value"], "size": 54, "color": BRAND, "font": FONT_XBOLD, "line": 1.0},
                    {"text": fact["label"], "size": 17, "color": TEXT2, "before": 6},
                ],
            )
    add_notes(slide, spec["notes"])


def _points(slide, points, x, y, w, gap=0.38):
    py = y
    for point in points:
        paragraphs = [
            {"text": point["lead"], "size": 21, "color": TEXT1, "font": FONT_SEMI, "line": 1.0},
            {"text": point["text"], "size": 16, "color": TEXT2, "before": 5, "line": 1.1},
        ]
        h = text_height(paragraphs, w - 0.27)
        add_box(slide, x, py + 0.05, 0.07, 0.3, fill=BRAND, shape=MSO_SHAPE.RECTANGLE, radius=None)
        add_text(slide, x + 0.27, py, w - 0.27, h + 0.05, paragraphs)
        py += h + gap
    return py - gap


def _flow(slide, items, x, y, w):
    pill_h = 0.44
    arrow_h = 0.27
    for i, item in enumerate(items):
        py = y + i * (pill_h + arrow_h)
        last = i == len(items) - 1
        middle = i == 1
        box = add_box(
            slide,
            x,
            py,
            w,
            pill_h,
            fill=BRAND_DIM if middle else BG3,
            line=BRAND if middle else LINE,
            radius=0.25,
        )
        box.name = "Поток"
        add_text(
            slide,
            x,
            py,
            w,
            pill_h,
            [{"text": item, "size": 16, "color": BRAND if middle else TEXT1, "font": FONT_SEMI, "align": PP_ALIGN.CENTER}],
            anchor=MSO_ANCHOR.MIDDLE,
        )
        if not last:
            add_box(slide, x + w / 2 - 0.13, py + pill_h + 0.05, 0.26, arrow_h - 0.1, fill=TEXT3, shape=MSO_SHAPE.DOWN_ARROW, radius=None)
    return y + len(items) * (pill_h + arrow_h) - arrow_h


def _stats(slide, stats, x, y, w, gap=0.42):
    py = y
    for stat in stats:
        paragraphs = [
            {"text": stat["value"], "size": 44, "color": BRAND, "font": FONT_XBOLD, "line": 1.0},
            {"text": stat["label"], "size": 17, "color": TEXT2, "before": 2, "line": 1.1},
        ]
        h = text_height(paragraphs, w)
        add_text(slide, x, py, w, h + 0.05, paragraphs)
        py += h + gap
    return py - gap


def _check_bottom(y: float, what: str) -> None:
    if y > SLIDE_H - 0.35:
        WARNINGS.append(f"слайд {CURRENT['slide']}: {what} заканчивается на {y:.2f} дюйма — ниже поля")


def layout_screen(prs, spec, number, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_header(slide, spec, number, total)
    col_x, col_w = MARGIN, 3.5
    img_x = MARGIN + col_w + 0.35
    img_w = SLIDE_W - 0.45 - img_x
    img_y, img_h = CONTENT_TOP, SLIDE_H - CONTENT_TOP - 0.42
    y = CONTENT_TOP + 0.05
    if spec.get("stats"):
        y = _stats(slide, spec["stats"], col_x, y - 0.12, col_w)
    if spec.get("points"):
        y = _points(slide, spec["points"], col_x, y, col_w)
    if spec.get("flow"):
        y = _flow(slide, spec["flow"], col_x, y + 0.36, col_w - 0.2)
    _check_bottom(y, "левая колонка")
    image = HERE / spec["image"]
    crop = spec.get("crop", [0, 0, 0, 0])
    if image.exists():
        x, yy, w, h = fit(image_aspect(image, crop), img_x, img_y, img_w, img_h, halign="center", valign="top")
        add_picture(slide, image, x, yy, w, h, crop=crop)
    else:
        add_box(slide, img_x, img_y, img_w, img_w * 9 / 16, fill=BG3, line=LINE)
        add_text(slide, img_x, img_y, img_w, img_w * 9 / 16, [{"text": f"нет файла {spec['image']}", "size": 18, "color": TEXT3, "align": PP_ALIGN.CENTER}], anchor=MSO_ANCHOR.MIDDLE)
    add_notes(slide, spec["notes"])


def layout_final(prs, spec, number, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, BG0)
    size = 1.3
    slide.shapes.add_picture(str(ASSETS / "logo.png"), Inches((SLIDE_W - size) / 2), Inches(1.35), Inches(size), Inches(size))
    add_text(slide, 1.0, 2.95, SLIDE_W - 2.0, 1.4, [{"text": spec["title"], "size": 80, "color": TEXT1, "font": FONT_XBOLD, "align": PP_ALIGN.CENTER, "line": 1.0}])
    add_text(slide, 1.0, 4.4, SLIDE_W - 2.0, 0.6, [{"text": spec["subtitle"], "size": 28, "color": BRAND, "font": FONT_MEDIUM, "align": PP_ALIGN.CENTER}])
    add_text(slide, 1.0, 5.6, SLIDE_W - 2.0, 0.5, [{"text": spec["meta"], "size": 18, "color": TEXT2, "align": PP_ALIGN.CENTER}])
    slide.shapes.add_picture(str(ASSETS / "hazard.png"), 0, Inches(SLIDE_H - 0.14), Inches(SLIDE_W), Inches(0.14))
    add_notes(slide, spec["notes"])


LAYOUTS = {"title": layout_title, "scheme": layout_scheme, "screen": layout_screen, "final": layout_final}


# ---------------------------------------------------------------- сборка


def build(spec_path: Path, out: Path) -> Path:
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    ensure_assets()
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    slides = spec["slides"]
    for i, slide_spec in enumerate(slides, 1):
        CURRENT["slide"] = i
        LAYOUTS[slide_spec["layout"]](prs, slide_spec, i, len(slides))
    props = prs.core_properties
    props.title = spec["meta"]["title"]
    props.author = spec["meta"]["author"]
    props.last_modified_by = spec["meta"]["author"]
    props.comments = ""
    props.keywords = ""
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    print("PPTX:", out, f"{out.stat().st_size / 1e6:.1f} МБ, слайдов: {len(slides)}")
    for warning in WARNINGS:
        print("Вёрстка:", warning)
    if not WARNINGS:
        print("Вёрстка: переполнений не найдено")
    return out


def to_pdf(pptx: Path) -> Path:
    soffice = shutil.which("soffice") or "/opt/homebrew/bin/soffice"
    profile = Path("/tmp/zz-deck-v2-lo-profile")
    cmd = [
        soffice,
        f"-env:UserInstallation=file://{profile}",
        "--headless",
        "--convert-to",
        "pdf",
        "--outdir",
        str(pptx.parent),
        str(pptx),
    ]
    # LibreOffice из Homebrew ищет шрифты через fontconfig без пользовательских каталогов:
    # без fonts.conf вместо Onest в PDF попадает Linux Libertine.
    env = dict(os.environ, FONTCONFIG_FILE=str(HERE / "fonts.conf"))
    subprocess.run(cmd, check=True, timeout=300, env=env)
    pdf = pptx.with_suffix(".pdf")
    fonts = subprocess.run(["pdffonts", str(pdf)], capture_output=True, text=True).stdout if shutil.which("pdffonts") else ""
    if fonts and "Onest" not in fonts:
        print("ВНИМАНИЕ: в PDF нет Onest — проверьте установку шрифта и fonts.conf")
    print("PDF:", pdf, f"{pdf.stat().st_size / 1e6:.1f} МБ")
    return pdf


def to_png(pdf: Path, out_dir: Path, dpi: int = 72) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)
    pdftoppm = shutil.which("pdftoppm") or "/opt/homebrew/bin/pdftoppm"
    subprocess.run([pdftoppm, "-r", str(dpi), "-png", str(pdf), str(out_dir / "slide")], check=True, timeout=300)
    print("PNG:", out_dir)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--spec", type=Path, default=HERE / "slides.json")
    parser.add_argument("--out", type=Path, default=None, help="путь PPTX (по умолчанию рядом со slides.json)")
    parser.add_argument("--pdf", action="store_true", help="сделать PDF через LibreOffice")
    parser.add_argument("--png", type=Path, default=None, help="каталог для PNG слайдов (нужен --pdf)")
    parser.add_argument("--dpi", type=int, default=72)
    args = parser.parse_args()
    spec = json.loads(args.spec.read_text(encoding="utf-8"))
    out = args.out or (HERE / f"{spec['meta']['output']}.pptx")
    pptx = build(args.spec, out)
    if args.pdf or args.png:
        pdf = to_pdf(pptx)
        if args.png:
            to_png(pdf, args.png, args.dpi)


if __name__ == "__main__":
    main()
